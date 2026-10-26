#!/usr/bin/env python3
"""The card spreadsheet (DECISIONS #43): src/data/cards.json <-> an .xlsx the designer edits.

  python3 scripts/cards-xlsx.py export [docs/cards.xlsx]   write the workbook from cards.json
  python3 scripts/cards-xlsx.py import [docs/cards.xlsx]   read an edited workbook back into cards.json
  python3 scripts/cards-xlsx.py check  [docs/cards.xlsx]   import to a temp file and diff against cards.json

Sheets: "Warrior cards" (basics, Bulwark, Berserker, in scope), "Other classes" (wizard and
shaman, out of scope, kept so the game data stays whole), "Skill tree" (nodes, costs,
prerequisites), "How to edit" (the effect syntax). Each effect is one cell in a small language
that round-trips exactly, e.g. "damage 100% +4", "magic 150%", "self guard 1",
"enemy burn 2 for 2 rounds", "enemy chill 65% chance", "heal 25%". After an import run
`npm run check`; the data tests reject a sheet that breaks a rule (unknown id, bad family, etc).
"""
import json, os, re, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
JSON_PATH = os.path.join(ROOT, 'src', 'data', 'cards.json')
DEFAULT_XLSX = os.path.join(ROOT, 'docs', 'cards.xlsx')

STATUSES = ['exposed', 'guard', 'ward', 'aegis', 'burn', 'chill', 'weak', 'stun', 'brace', 'scatter', 'curse', 'harden']
RARITIES = ['basic', 'common', 'uncommon', 'rare']
USES = ['', 'fleeting', 'once per fight']
CLASS_ORDER = ['warrior', 'wizard', 'shaman']
CARD_COLUMNS = ['id', 'name', 'class', 'family', 'rarity', 'uses', 'effect 1', 'effect 2', 'effect 3', 'text (optional)', 'notes']

# ------------------------------------------------------------------ effect <-> text

def pct(x):
    v = round(x * 100, 4)
    return f"{int(v)}%" if v == int(v) else f"{v}%"

def effect_to_text(e):
    t = e['type']
    if t == 'damage':
        parts = ['magic' if e.get('kind') == 'magic' else 'damage', pct(e['mult'])]
        if e.get('flat'): parts.append(f"+{e['flat']}")
        if e.get('hits', 1) != 1: parts.append(f"x{e['hits']}")
        if e.get('ignoreDef'): parts.append('ignore def')
        if e.get('lifesteal'): parts.append(f"drain {pct(e['lifesteal'])}")
        return ' '.join(parts)
    if t == 'heal':
        if e.get('pct'): return f"heal {pct(e['pct'])}" + (f" +{e['flat']}" if e.get('flat') else '')
        return f"heal {e.get('flat', 0)}"
    if t == 'status':
        st = e['status']
        parts = [e['target'], st]
        if st in ('guard', 'ward'):
            parts.append(str(e.get('magnitude', 1)))                       # hits blocked; turns are always 99
        else:
            if e.get('magnitude') is not None: parts.append(str(e['magnitude']))
            if st != 'chill': parts.append(f"for {e['turns']} round" + ('' if e['turns'] == 1 else 's'))
        if e.get('chance') is not None: parts.append(f"{pct(e['chance'])} chance")
        if e.get('kind') == 'magic': parts.append('magic')
        return ' '.join(parts)
    raise ValueError(f"unknown effect type {t}")

def frac(tok):
    """'150%' -> 1.5, '100%' -> 1 (an integral value stays an int, as the game data writes it)."""
    v = float(tok.rstrip('%')) / 100
    return int(v) if v == int(v) else round(v, 6)

def text_to_effect(text, where):
    s = ' '.join(text.strip().lower().split())
    if not s: return None
    def fail(msg): raise ValueError(f"{where}: {msg} in '{text}'")
    toks = s.split(' ')
    head = toks[0]
    if head in ('damage', 'magic'):
        m = re.fullmatch(r'(\d+(?:\.\d+)?)%', toks[1] if len(toks) > 1 else '') or fail('expected a percentage after ' + head)
        e = {'type': 'damage', 'mult': frac(m.group(1))}
        if head == 'magic': e['kind'] = 'magic'
        else: e['kind'] = 'physical'
        i = 2
        while i < len(toks):
            tk = toks[i]
            if re.fullmatch(r'\+\d+', tk): e['flat'] = int(tk[1:]); i += 1
            elif re.fullmatch(r'x\d+', tk): e['hits'] = int(tk[1:]); i += 1
            elif tk == 'ignore' and i + 1 < len(toks) and toks[i + 1] == 'def': e['ignoreDef'] = True; i += 2
            elif tk == 'drain' and i + 1 < len(toks) and re.fullmatch(r'\d+%', toks[i + 1]): e['lifesteal'] = frac(toks[i + 1]); i += 2
            else: fail(f"unexpected '{tk}'")
        if e['kind'] == 'physical': del e['kind']
        return tidy(e)
    if head == 'heal':
        e = {'type': 'heal'}
        for tk in toks[1:]:
            if re.fullmatch(r'\d+%', tk): e['pct'] = frac(tk)
            elif re.fullmatch(r'\+?\d+', tk): e['flat'] = int(tk.lstrip('+'))
            else: fail(f"unexpected '{tk}'")
        if 'pct' not in e and 'flat' not in e: fail('heal needs a % or a flat amount')
        return tidy(e)
    if head in ('self', 'enemy'):
        if len(toks) < 2 or toks[1] not in STATUSES: fail('expected a status after ' + head + ': ' + ', '.join(STATUSES))
        st = toks[1]
        e = {'type': 'status', 'target': head, 'status': st}
        i = 2
        if i < len(toks) and re.fullmatch(r'\d+', toks[i]): e['magnitude'] = int(toks[i]); i += 1
        while i < len(toks):
            tk = toks[i]
            if tk == 'for' and i + 2 < len(toks) and re.fullmatch(r'\d+', toks[i + 1]) and toks[i + 2] in ('round', 'rounds'): e['turns'] = int(toks[i + 1]); i += 3
            elif re.fullmatch(r'\d+%', tk) and i + 1 < len(toks) and toks[i + 1] == 'chance': e['chance'] = frac(tk); i += 2
            elif tk == 'magic': e['kind'] = 'magic'; i += 1
            else: fail(f"unexpected '{tk}'")
        if st in ('guard', 'ward'):
            e['turns'] = 99
            e.setdefault('magnitude', 1)
        elif st == 'chill':
            e.setdefault('turns', 1)
        if 'turns' not in e: fail(f"{st} needs 'for N rounds'")
        return tidy(e)
    fail("an effect starts with damage, magic, heal, self or enemy")

def tidy(e):
    """Key order the game data uses, so a round trip reproduces cards.json byte for byte."""
    order = ['type', 'target', 'status', 'mult', 'flat', 'hits', 'ignoreDef', 'kind', 'lifesteal', 'pct', 'turns', 'magnitude', 'chance']
    out = {k: e[k] for k in order if k in e}
    for k in e:
        if k not in out: out[k] = e[k]
    return out

# ------------------------------------------------------------------ rows

def uses_out(u): return 'once per fight' if u == 'once_per_fight' else (u or '')
def uses_in(u):
    u = (u or '').strip().lower()
    if u in ('', 'none'): return None
    if u == 'fleeting': return 'fleeting'
    if u in ('once per fight', 'once_per_fight'): return 'once_per_fight'
    raise ValueError(f"uses must be blank, fleeting or once per fight, not '{u}'")

def family_name_of(data, class_id, skill_id):
    for st in data['sets']:
        if st['classId'] == class_id and skill_id in st['skills']: return st['name']
    return ''

def card_rows(data, class_ids):
    rows = []
    for sk in data['skills']:
        if sk['rarity'] == 'basic':
            if 'warrior' not in class_ids: continue
            fam = 'basic'
        elif sk['classId'] == 'any':
            # General tomes are listed once, on the warrior sheet, with their warrior family; the other
            # classes keep whatever family held them (out of scope).
            if 'warrior' not in class_ids: continue
            fam = family_name_of(data, 'warrior', sk['id']) or 'none'
        else:
            if sk['classId'] not in class_ids: continue
            fam = family_name_of(data, sk['classId'], sk['id']) or 'none'
        effects = [effect_to_text(e) for e in sk['effects']]
        if len(effects) > 3: raise ValueError(f"{sk['id']}: more than 3 effects; widen the sheet")
        effects += [''] * (3 - len(effects))
        rows.append([sk['id'], sk['name'], sk['classId'], fam, sk['rarity'], uses_out(sk.get('uses')), *effects, sk.get('text', ''), sk.get('note', '')])
    return rows

# ------------------------------------------------------------------ export

HEAD_FILL = PatternFill('solid', fgColor='1F2A44')
EDIT_FILL = PatternFill('solid', fgColor='FFF7CC')
LOCK_FILL = PatternFill('solid', fgColor='EDEDED')
FONT = 'Arial'

def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF')
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.freeze_panes = 'B2'

def write_cards_sheet(ws, rows, families, editable=True):
    ws.append(CARD_COLUMNS)
    for r in rows: ws.append(r)
    style_header(ws, len(CARD_COLUMNS))
    widths = [16, 20, 9, 24, 11, 15, 30, 30, 30, 48, 48]
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.fill = EDIT_FILL if (editable and cell.column > 1) else LOCK_FILL
    if editable and ws.max_row > 1:
        n = ws.max_row
        dv_fam = DataValidation(type='list', formula1='"' + ','.join(['basic', 'none'] + families) + '"', allow_blank=False)
        dv_rar = DataValidation(type='list', formula1='"' + ','.join(RARITIES) + '"', allow_blank=False)
        dv_use = DataValidation(type='list', formula1='"' + ','.join(USES[1:]) + '"', allow_blank=True)
        for dv, col in ((dv_fam, 'D'), (dv_rar, 'E'), (dv_use, 'F')):
            dv.error = 'Pick one of the listed values.'; dv.errorTitle = 'Not a valid value'
            ws.add_data_validation(dv)
            dv.add(f'{col}2:{col}{n}')

def export(path):
    data = json.load(open(JSON_PATH))
    wb = Workbook()
    ws = wb.active; ws.title = 'Warrior cards'
    warrior_fams = [st['name'] for st in data['sets'] if st['classId'] == 'warrior']
    write_cards_sheet(ws, card_rows(data, ['warrior']), warrior_fams, editable=True)
    ws2 = wb.create_sheet('Other classes')
    other_fams = [st['name'] for st in data['sets'] if st['classId'] != 'warrior']
    write_cards_sheet(ws2, card_rows(data, ['wizard', 'shaman']), other_fams, editable=False)
    ws2['A1'].comment = None
    ws3 = wb.create_sheet('Skill tree')
    ws3.append(['class', 'node id', 'family name', 'cost (points)', 'requires (node id, blank if none)', 'cards in this family (from the cards sheets; read only)'])
    for c in CLASS_ORDER:
        for n in data['trees'][c]:
            st = next(s for s in data['sets'] if s['id'] == n['setId'])
            ws3.append([c, n['setId'], st['name'], n['cost'], ', '.join(n.get('requires', [])), ', '.join(st['skills'])])
    style_header(ws3, 6)
    for i, w in enumerate([10, 16, 16, 14, 30, 70], 1): ws3.column_dimensions[get_column_letter(i)].width = w
    for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
        for cell in row:
            cell.font = Font(name=FONT); cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.fill = EDIT_FILL if (row[0].value == 'warrior' and cell.column in (3, 4, 5)) else LOCK_FILL
    ws4 = wb.create_sheet('How to edit')
    guide = [
        ['Re-RPG card sheet', ''],
        ['', ''],
        ['What to edit', 'Yellow cells on "Warrior cards" and "Skill tree". Grey cells are read only (ids, and the other two classes, which are out of scope). Keep the id column as it is; it is how the game and the save files know a card.'],
        ['Where a card sits', 'The "family" column is the skill-tree node that unlocks the card: basic (in every warrior fight), Bulwark, Berserker, or none (exists but never drops). To move a card between nodes, change this one cell.'],
        ['What a card does', 'Up to three effects, one per cell, in order. The game computes the real numbers from the hero\'s attack, so a percentage here is of class attack; gear attack is added flat to every hit.'],
        ['', ''],
        ['Effect syntax', 'Example'],
        ['damage N%', 'damage 100%   (a plain hit)'],
        ['damage N% +F', 'damage 100% +4   (flat bonus after the percentage)'],
        ['damage N% xH', 'damage 60% x3   (H hits, each N%)'],
        ['damage N% ignore def', 'damage 160% ignore def'],
        ['damage N% drain P%', 'damage 120% drain 50%   (heals P% of the damage dealt)'],
        ['magic N% ...', 'magic 150%   (same options; Ward blocks it instead of Guard)'],
        ['heal N%', 'heal 25%   (of max HP)   or   heal 5   (flat)'],
        ['self guard H', 'self guard 1   (block the next H physical hits)'],
        ['self ward H', 'self ward 1   (block the next H magic hits or curses)'],
        ['self aegis for N rounds', 'self aegis for 2 rounds   (block everything)'],
        ['self STATUS for N rounds', 'self exposed for 1 round   or   self weak for 2 rounds'],
        ['enemy STATUS for N rounds', 'enemy exposed for 2 rounds'],
        ['enemy burn M for N rounds', 'enemy burn 2 for 2 rounds   (M damage per round; curse works the same)'],
        ['enemy chill P% chance', 'enemy chill 65% chance   (skips its next action)'],
        ['... P% chance', 'any status may end with "P% chance"; "magic" after a hostile status means Ward can block it'],
        ['Statuses', ', '.join(STATUSES)],
        ['', ''],
        ['Timing', 'A status ticks at the end of the round it was applied in. "self weak for 2 rounds" covers this round and your next turn; "enemy exposed for 2 rounds" covers your next turn; "self exposed for 1 round" covers only this round\'s enemy attacks. See DECISIONS #40.'],
        ['uses', 'blank (a normal tome), fleeting (destroyed when played), once per fight (back next fight).'],
        ['rarity', 'basic cards are in every fight and never drop; common / uncommon / rare set how often a card is offered. Only the basics should be basic.'],
        ['text (optional)', 'Leave blank and the import writes a plain description from the effects. The card in the game always shows computed numbers anyway.'],
        ['Skill tree sheet', 'cost is skill points (one per level). requires is another node id that must be owned first. Family names can be renamed here.'],
        ['Sending it back', 'Save as .xlsx and send the file. It is imported with: python3 scripts/cards-xlsx.py import docs/cards.xlsx, then npm run check.'],
    ]
    for r in guide: ws4.append(r)
    ws4.column_dimensions['A'].width = 30; ws4.column_dimensions['B'].width = 110
    for row in ws4.iter_rows(min_row=1, max_row=ws4.max_row):
        for cell in row:
            cell.font = Font(name=FONT, bold=(cell.column == 1 or cell.row == 1), size=14 if cell.row == 1 else 11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    wb.save(path)
    print(f"wrote {path}: {len(card_rows(data, ['warrior']))} warrior cards, {len(card_rows(data, ['wizard', 'shaman']))} other, {sum(len(v) for v in data['trees'].values())} tree nodes")

# ------------------------------------------------------------------ import

def describe(effects, uses):
    """A plain text for cards whose text cell was left blank (the game shows computed numbers anyway)."""
    bits = []
    for e in effects:
        if e['type'] == 'damage':
            s = f"Deal {pct(e['mult'])}" + (f" + {e['flat']}" if e.get('flat') else '') + (' magic' if e.get('kind') == 'magic' else '') + ' damage'
            if e.get('hits', 1) > 1: s = f"{e['hits']} hits of {pct(e['mult'])}" + (' magic' if e.get('kind') == 'magic' else '') + ' damage'
            if e.get('ignoreDef'): s += ', ignoring defence'
            if e.get('lifesteal'): s += f", heal {pct(e['lifesteal'])} of it"
            bits.append(s + '.')
        elif e['type'] == 'heal':
            bits.append('Heal ' + (f"{pct(e['pct'])} HP" if e.get('pct') else f"{e.get('flat', 0)} HP") + '.')
        else:
            st = e['status']; who = 'You are' if e['target'] == 'self' else 'Enemy is'
            if st in ('guard', 'ward'): bits.append(f"Block the next {e.get('magnitude', 1)} {'physical' if st == 'guard' else 'magic'} hit{'s' if e.get('magnitude', 1) != 1 else ''}.")
            elif st == 'aegis': bits.append(f"Block every hit for {e['turns']} rounds.")
            elif st in ('burn', 'curse'): bits.append(f"{st.title()} {e.get('magnitude', 1)} for {e['turns']} rounds.")
            elif st == 'chill': bits.append((f"{pct(e['chance'])} chance to " if e.get('chance') else '') + 'Chill: the enemy skips its next action.')
            else: bits.append(f"{who} {st.title()} for {e['turns']} round{'s' if e['turns'] != 1 else ''}.")
    if uses == 'once_per_fight': bits.append('Once per fight.')
    if uses == 'fleeting': bits.append('Fleeting: spent when played.')
    return ' '.join(bits)

def read_cards_sheet(ws, data, class_ids, fam_by_name, skills_out, sets_skills):
    header = [c.value for c in ws[1]]
    if header[:9] != CARD_COLUMNS[:9]: raise ValueError(f"{ws.title}: columns changed; expected {CARD_COLUMNS[:9]}")
    known = {s['id']: s for s in data['skills']}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]: continue
        sid, name, cls, fam, rarity, uses, e1, e2, e3, text, note = (list(row) + [None] * 11)[:11]
        sid = str(sid).strip(); where = f"{ws.title} / {sid}"
        if sid not in known: raise ValueError(f"{where}: unknown card id (ids are read only; add new cards in cards.json first)")
        cls = (cls or known[sid]['classId']).strip()
        if cls not in class_ids and cls != 'any': raise ValueError(f"{where}: class {cls} does not belong on this sheet")
        rarity = (rarity or '').strip().lower()
        if rarity not in RARITIES: raise ValueError(f"{where}: rarity must be one of {RARITIES}")
        effects = [x for x in (text_to_effect(str(c), where) if c not in (None, '') else None for c in (e1, e2, e3)) if x]
        if not effects: raise ValueError(f"{where}: a card needs at least one effect")
        u = uses_in(uses)
        sk = {'id': sid, 'name': str(name).strip(), 'text': (str(text).strip() if text not in (None, '') else describe(effects, u)), 'classId': cls}
        fam_s = str(fam or '').strip()
        if rarity == 'basic':
            if fam_s not in ('', 'basic'): raise ValueError(f"{where}: a basic card's family is 'basic'")
        elif cls == 'any':
            # the warrior family for this general tome; blank or none means no warrior family drops it
            if fam_s.lower() not in ('none', ''):
                if ('warrior', fam_s) not in fam_by_name: raise ValueError(f"{where}: no warrior family '{fam_s}'")
                sets_skills[fam_by_name[('warrior', fam_s)]].append(sid)
        else:
            if fam_s.lower() not in ('none', ''):
                if (cls, fam_s) not in fam_by_name: raise ValueError(f"{where}: no family '{fam_s}' for {cls}; families are {[f for (c, f) in fam_by_name if c == cls]}")
                set_id = fam_by_name[(cls, fam_s)]
                sets_skills[set_id].append(sid)
                sk['setId'] = set_id
        sk['rarity'] = rarity
        if u: sk['uses'] = u
        sk['effects'] = effects
        if note not in (None, ''): sk['note'] = str(note).strip()
        skills_out[sid] = sk

def do_import(path, out_path=JSON_PATH):
    data = json.load(open(JSON_PATH))
    wb = load_workbook(path)
    # Skill tree first: it may rename families and change costs.
    ws3 = wb['Skill tree']
    trees = {c: [] for c in CLASS_ORDER}
    sets_by_id = {s['id']: dict(s) for s in data['sets']}
    for row in ws3.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]: continue
        cls, node_id, fam_name, cost, requires = (list(row) + [None] * 5)[:5]
        cls = str(cls).strip(); node_id = str(node_id).strip()
        if node_id not in sets_by_id: raise ValueError(f"Skill tree: unknown node id {node_id}")
        sets_by_id[node_id]['name'] = str(fam_name).strip()
        node = {'setId': node_id, 'cost': int(cost)}
        req = [r.strip() for r in str(requires or '').split(',') if r.strip()]
        for r in req:
            if r not in sets_by_id: raise ValueError(f"Skill tree: {node_id} requires unknown node {r}")
        if req: node['requires'] = req
        trees[cls].append(node)
    # A family renamed on the tree sheet is still found by its old name on the card sheets.
    fam_by_name = {(s['classId'], s['name']): s['id'] for s in sets_by_id.values()}
    for s in data['sets']: fam_by_name.setdefault((s['classId'], s['name']), s['id'])
    # Non-warrior families keep the general tomes they already held (those classes are out of scope).
    general = {s['id'] for s in data['skills'] if s['classId'] == 'any' and s['rarity'] != 'basic'}
    sets_skills = {s['id']: ([i for i in s['skills'] if i in general] if s['classId'] != 'warrior' else []) for s in data['sets']}
    skills_out = {}
    read_cards_sheet(wb['Warrior cards'], data, ['warrior'], fam_by_name, skills_out, sets_skills)
    read_cards_sheet(wb['Other classes'], data, ['wizard', 'shaman'], fam_by_name, skills_out, sets_skills)
    # Keep the original card order; a card missing from both sheets is an error, not a deletion.
    missing = [s['id'] for s in data['skills'] if s['id'] not in skills_out]
    if missing: raise ValueError(f"cards missing from the workbook: {missing}")
    skills = [skills_out[s['id']] for s in data['skills']]
    sets = []
    for s in data['sets']:
        st = sets_by_id[s['id']]
        # class cards first in card order, then the general tomes in the order the family already listed them (new ones last)
        order = {sk['id']: i for i, sk in enumerate(skills)}
        prev = {i: n for n, i in enumerate(s['skills'])}
        members = list(dict.fromkeys(sets_skills[st['id']]))
        sets.append({'id': st['id'], 'classId': st['classId'], 'name': st['name'], 'skills': sorted(members, key=lambda i: (skills[order[i]]['classId'] == 'any', prev.get(i, 10 ** 6) if skills[order[i]]['classId'] == 'any' else order[i]))})
    out = {'basics': data['basics'], 'skills': skills, 'sets': sets, 'trees': trees}
    with open(out_path, 'w') as f: json.dump(out, f, indent=2); f.write('\n')
    print(f"wrote {out_path} from {path}")

if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else 'export'
    path = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_XLSX
    if cmd == 'export': export(path)
    elif cmd == 'import': do_import(path)
    elif cmd == 'check':
        tmp = path + '.check.json'
        do_import(path, tmp)
        same = open(tmp).read() == open(JSON_PATH).read()
        os.remove(tmp)
        print('round trip:', 'identical' if same else 'DIFFERS')
        sys.exit(0 if same else 1)
    else:
        print(__doc__); sys.exit(2)
